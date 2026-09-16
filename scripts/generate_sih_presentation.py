import sys, os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

sys.stdout.reconfigure(encoding='utf-8')

prs = Presentation()
# Set widescreen 16:9 (13.333 x 7.5 inches)
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Color Palette
DARK_NAVY = RGBColor(9, 13, 22)
HEADER_BLUE = RGBColor(16, 44, 87)
ACCENT_CYAN = RGBColor(2, 132, 199)
ACCENT_GREEN = RGBColor(16, 185, 129)
ACCENT_AMBER = RGBColor(217, 119, 6)
SLATE_TEXT = RGBColor(30, 41, 59)
MUTED_SLATE = RGBColor(100, 116, 139)
WHITE = RGBColor(255, 255, 255)
CARD_BG = RGBColor(241, 245, 249)
BORDER_COLOR = RGBColor(203, 213, 225)
BLUE_BAR = RGBColor(0, 102, 178)

blank_layout = prs.slide_layouts[6]

def add_header(slide, title_text, slide_num):
    # Top left team circle
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.4), Inches(0.25), Inches(1.5), Inches(0.75))
    circle.fill.solid()
    circle.fill.fore_color.rgb = WHITE
    circle.line.color.rgb = RGBColor(100, 116, 139)
    circle.line.width = Pt(1.5)
    tf_c = circle.text_frame
    tf_c.word_wrap = True
    p_c = tf_c.paragraphs[0]
    p_c.text = "CODEZILLA"
    p_c.font.size = Pt(11)
    p_c.font.bold = True
    p_c.font.color.rgb = SLATE_TEXT
    p_c.alignment = PP_ALIGN.CENTER

    # Title
    tb = slide.shapes.add_textbox(Inches(2.2), Inches(0.2), Inches(8.5), Inches(0.9))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.name = "Georgia"
    p.font.color.rgb = HEADER_BLUE
    p.alignment = PP_ALIGN.CENTER

    # Bottom blue bar
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.05), Inches(13.333), Inches(0.45))
    bar.fill.solid()
    bar.fill.fore_color.rgb = BLUE_BAR
    bar.line.fill.background()
    tf_b = bar.text_frame
    p_b = tf_b.paragraphs[0]
    p_b.text = f"@SIH Idea submission- Template                                                                                      {slide_num}"
    p_b.font.size = Pt(11)
    p_b.font.color.rgb = WHITE

# ==========================================
# SLIDE 1: TITLE PAGE
# ==========================================
slide1 = prs.slides.add_slide(blank_layout)

# SIH Header
tb_sih = slide1.shapes.add_textbox(Inches(1.0), Inches(0.4), Inches(11.333), Inches(0.8))
p_sih = tb_sih.text_frame.paragraphs[0]
p_sih.text = "SMART INDIA HACKATHON 2026"
p_sih.font.size = Pt(34)
p_sih.font.bold = True
p_sih.font.name = "Georgia"
p_sih.font.color.rgb = HEADER_BLUE
p_sih.alignment = PP_ALIGN.CENTER

tb_t = slide1.shapes.add_textbox(Inches(1.0), Inches(1.2), Inches(11.333), Inches(0.6))
p_t = tb_t.text_frame.paragraphs[0]
p_t.text = "TITLE PAGE"
p_t.font.size = Pt(24)
p_t.font.bold = True
p_t.font.name = "Georgia"
p_t.font.color.rgb = SLATE_TEXT
p_t.alignment = PP_ALIGN.CENTER

# Details Card
card1 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(2.0), Inches(11.333), Inches(4.7))
card1.fill.solid()
card1.fill.fore_color.rgb = CARD_BG
card1.line.color.rgb = BORDER_COLOR
card1.line.width = Pt(1.5)

tf1 = card1.text_frame
tf1.word_wrap = True
tf1.margin_left = Inches(0.5)
tf1.margin_top = Inches(0.35)

items = [
    ("Problem Statement ID", "SIH26-26145"),
    ("Problem Statement Title", "AI-Based Detection of Cyber Threats in Unidirectional IP Traffic"),
    ("Organization / Ministry", "National Technical Research Organisation (NTRO)"),
    ("Theme", "Blockchain & Cybersecurity"),
    ("PS Category", "Software"),
    ("Team ID", "AMITY-2026-899F3DA7"),
    ("Team Name", "CODEZILLA"),
    ("Project Title", "CyberSentinel: Real-Time Passive Threat Intelligence & AI Detection")
]

for idx, (label, val) in enumerate(items):
    p = tf1.add_paragraph() if idx > 0 else tf1.paragraphs[0]
    p.space_after = Pt(8)
    run1 = p.add_run()
    run1.text = f"•  {label} — "
    run1.font.bold = True
    run1.font.size = Pt(17)
    run1.font.color.rgb = SLATE_TEXT
    run2 = p.add_run()
    run2.text = val
    run2.font.bold = (label in ["Problem Statement ID", "Team Name", "Project Title"])
    run2.font.size = Pt(17)
    run2.font.color.rgb = ACCENT_CYAN if label in ["Problem Statement ID", "Team Name", "Project Title"] else SLATE_TEXT

# Bottom bar
bar1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.05), Inches(13.333), Inches(0.45))
bar1.fill.solid()
bar1.fill.fore_color.rgb = BLUE_BAR
bar1.line.fill.background()
tf_b1 = bar1.text_frame
p_b1 = tf_b1.paragraphs[0]
p_b1.text = "@SIH Idea submission- Template                                                                                      1"
p_b1.font.size = Pt(11)
p_b1.font.color.rgb = WHITE

# ==========================================
# SLIDE 2: IDEA TITLE & PROPOSED SOLUTION
# ==========================================
slide2 = prs.slides.add_slide(blank_layout)
add_header(slide2, "CYBERSENTINEL — PROPOSED SOLUTION", 2)

# Main container
sub_tb = slide2.shapes.add_textbox(Inches(0.8), Inches(1.15), Inches(11.7), Inches(0.5))
p_sub = sub_tb.text_frame.paragraphs[0]
p_sub.text = "❖ Proposed Solution (Describe your Idea/Solution/Prototype)"
p_sub.font.size = Pt(20)
p_sub.font.bold = True
p_sub.font.color.rgb = BLUE_BAR

# 3 Pillars
col_w = Inches(3.75)
gap = Inches(0.2)
top_pos = Inches(1.75)
h_pos = Inches(5.05)

pillar_data = [
    ("1. Detailed Solution Architecture", [
        "100% Read-Only Passive Ingestion: Operates downstream of a hardware optical diode / data diode tap with zero packet transmission.",
        "Sub-Millisecond Flow Aggregator: Assembles unidirectional flows (TCP/UDP/ICMP) without needing return handshake (SYN-ACK / FIN).",
        "Dual-Engine Detection: Rule-based heuristic multi-stage engines running concurrently with HistGradientBoosting ML models.",
        "Full 2026 Enterprise SOC UI: Complete live dashboard with Dark/Light modes, topology graph, and VirusTotal-style threat scanner."
    ]),
    ("2. How It Solves the NTRO Problem", [
        "Eliminates Asymmetric Blindspots: Statistical flow reconstruction decodes beaconing & exfiltration without bidirectional state.",
        "Zero Reverse-Channel Intrusion: Complete physical air-gap preservation; impossible for attackers to detect or compromise the sensor.",
        "Multi-Threat Coverage: Actively detects C2 Beaconing, SYN/UDP DDoS Floods, Port/Host Scanning, DNS Exfiltration, and Brute Force.",
        "Real-Time WebSocket Stream: Zero polling latency (<15ms) delivering flow telemetry and threat updates directly to analysts."
    ]),
    ("3. Innovation & Uniqueness", [
        "Explainable AI (XAI) Feature Contributions: Generates human-readable evidence bars (IAT variance, Shannon Entropy, Fan-out).",
        "Markov Attack Trajectory (DTMC): Discrete-Time Markov model predicts probable lateral progression (Recon → Scan → C2 → Exfil).",
        "Online Baseline Calibration: 5-minute passive learning window continuously models normal operational bandwidth & ports.",
        "VirusTotal Multi-Engine Scanner: Integrated file & URL threat scanner providing instant Shannon entropy & DGA heuristics."
    ])
]

for i, (p_title, bullets) in enumerate(pillar_data):
    left_pos = Inches(0.8) + i * (col_w + gap)
    card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, top_pos, col_w, h_pos)
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = BORDER_COLOR
    card.line.width = Pt(1.5)
    
    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.2)
    tf.margin_right = Inches(0.2)
    tf.margin_top = Inches(0.2)
    
    p = tf.paragraphs[0]
    p.text = p_title
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = HEADER_BLUE
    p.space_after = Pt(10)
    
    for b in bullets:
        pb = tf.add_paragraph()
        pb.space_after = Pt(8)
        parts = b.split(":", 1)
        r1 = pb.add_run()
        r1.text = "• " + parts[0] + (":" if len(parts) > 1 else "")
        r1.font.bold = True
        r1.font.size = Pt(12)
        r1.font.color.rgb = SLATE_TEXT
        if len(parts) > 1:
            r2 = pb.add_run()
            r2.text = parts[1]
            r2.font.size = Pt(11.5)
            r2.font.color.rgb = SLATE_TEXT

# ==========================================
# SLIDE 3: TECHNICAL APPROACH
# ==========================================
slide3 = prs.slides.add_slide(blank_layout)
add_header(slide3, "TECHNICAL APPROACH & ARCHITECTURE", 3)

# Subtitle
sub_tb3 = slide3.shapes.add_textbox(Inches(0.8), Inches(1.15), Inches(11.7), Inches(0.4))
p_sub3 = sub_tb3.text_frame.paragraphs[0]
p_sub3.text = "❖ Technologies Used & Implementation Methodology"
p_sub3.font.size = Pt(18)
p_sub3.font.bold = True
p_sub3.font.color.rgb = BLUE_BAR

# Left Box: Tech Stack Table / Pillars
w_left = Inches(5.6)
card_stack = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.65), w_left, Inches(5.15))
card_stack.fill.solid()
card_stack.fill.fore_color.rgb = CARD_BG
card_stack.line.color.rgb = BORDER_COLOR
card_stack.line.width = Pt(1.5)

tf_s = card_stack.text_frame
tf_s.word_wrap = True
tf_s.margin_left = Inches(0.25)
tf_s.margin_top = Inches(0.2)

p = tf_s.paragraphs[0]
p.text = "Technologies & Frameworks"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = HEADER_BLUE
p.space_after = Pt(10)

tech_specs = [
    ("Backend API Engine", "Python 3.14, FastAPI, Starlette, Uvicorn, SQLite (Async ORM)"),
    ("Machine Learning Core", "scikit-learn 1.7 (HistGradientBoosting Classifier), NumPy, Pandas"),
    ("Explainable AI (XAI)", "Inter-Arrival Time (IAT) Variance, Byte Entropy (Shannon), Fan-out Ratio"),
    ("Network Telemetry", "Zeek / Bro sensor integration, PCAP Replay engine, Raw Socket Diode Tap"),
    ("Enterprise SOC UI", "Modern Responsive Frontend, Vanilla JS, Chart.js 4.4, D3.js 7.0, Inter Font"),
    ("Real-Time Protocols", "Dual WebSockets (/ws/live for telemetry, /ws/alerts for incident push)"),
    ("Threat Intelligence", "MITRE ATT&CK Mapping (Recon, Discovery, C2, Exfil), VirusTotal Heuristics")
]

for label, desc in tech_specs:
    p = tf_s.add_paragraph()
    p.space_after = Pt(7)
    r1 = p.add_run()
    r1.text = f"• {label}: "
    r1.font.bold = True
    r1.font.size = Pt(12)
    r1.font.color.rgb = ACCENT_CYAN
    r2 = p.add_run()
    r2.text = desc
    r2.font.size = Pt(11.5)
    r2.font.color.rgb = SLATE_TEXT

# Right Box: Pipeline Methodology
w_right = Inches(5.85)
card_method = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.65), Inches(1.65), w_right, Inches(5.15))
card_method.fill.solid()
card_method.fill.fore_color.rgb = CARD_BG
card_method.line.color.rgb = BORDER_COLOR
card_method.line.width = Pt(1.5)

tf_m = card_method.text_frame
tf_m.word_wrap = True
tf_m.margin_left = Inches(0.25)
tf_m.margin_top = Inches(0.2)

p = tf_m.paragraphs[0]
p.text = "End-to-End Pipeline & Methodology"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = HEADER_BLUE
p.space_after = Pt(10)

method_steps = [
    ("Step 1: Passive Diode Ingestion", "Raw network packets captured via hardware tap with physical TX pins disconnected; 100% one-way ingress."),
    ("Step 2: Normalization & Aggregation", "Asymmetric IP flows grouped by (Src IP, Dst IP, Port, Proto); statistical feature vectors generated."),
    ("Step 3: Dual-Stage Detection", "6 Rule Detectors (C2, DDoS, Recon, Exfil, DNS, Auth) execute in parallel with HistGradientBoosting ML model."),
    ("Step 4: Risk Scoring & XAI", "Composite Risk Score (0-100) calculated with transparent SHAP-style feature contributions."),
    ("Step 5: DTMC Attack Trajectory", "Discrete-Time Markov Chain maps attack progression and computes transition probabilities to next state."),
    ("Step 6: Analyst SOC Command Center", "Broadcasts events to live SOC console with interactive D3 network graphs and non-intrusive containment playbooks.")
]

for label, desc in method_steps:
    p = tf_m.add_paragraph()
    p.space_after = Pt(7)
    r1 = p.add_run()
    r1.text = f"▶ {label}\n"
    r1.font.bold = True
    r1.font.size = Pt(12)
    r1.font.color.rgb = SLATE_TEXT
    r2 = p.add_run()
    r2.text = f"   {desc}"
    r2.font.size = Pt(11)
    r2.font.color.rgb = MUTED_SLATE

# ==========================================
# SLIDE 4: FEASIBILITY AND VIABILITY
# ==========================================
slide4 = prs.slides.add_slide(blank_layout)
add_header(slide4, "FEASIBILITY AND VIABILITY", 4)

sub_tb4 = slide4.shapes.add_textbox(Inches(0.8), Inches(1.15), Inches(11.7), Inches(0.4))
p_sub4 = sub_tb4.text_frame.paragraphs[0]
p_sub4.text = "❖ Technical Feasibility, Challenges & Mitigation Strategies"
p_sub4.font.size = Pt(18)
p_sub4.font.bold = True
p_sub4.font.color.rgb = BLUE_BAR

f_col_w = Inches(5.6)
# Feasibility Card
card_f = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.65), f_col_w, Inches(5.15))
card_f.fill.solid()
card_f.fill.fore_color.rgb = CARD_BG
card_f.line.color.rgb = BORDER_COLOR
card_f.line.width = Pt(1.5)

tf_f = card_f.text_frame
tf_f.word_wrap = True
tf_f.margin_left = Inches(0.25)
tf_f.margin_top = Inches(0.2)

p = tf_f.paragraphs[0]
p.text = "Analysis of Feasibility"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = HEADER_BLUE
p.space_after = Pt(10)

feas_points = [
    ("Working Production Prototype", "Already fully implemented, benchmarked, and verified end-to-end with 34 active FastAPI endpoints and sub-10ms inference latency."),
    ("Hardware Agnostic Deployment", "Runs on standard COTS Linux/Windows servers, industrial ruggedized PCs, or tactical air-gapped data diode appliances."),
    ("Zero Operational Overhead", "Requires zero endpoint agent installation, zero TLS private key sharing, and zero modifications to monitored enterprise network topology."),
    ("Resource Efficiency", "HistGradientBoosting classifier and statistical aggregators consume <150MB RAM and minimal CPU, easily handling >25,000 flows/sec."),
    ("Regulatory Compliance", "Full compliance with CERT-In, NTRO guidelines, ISO 27001, and Indian Digital Personal Data Protection (DPDP) Act.")
]

for label, desc in feas_points:
    p = tf_f.add_paragraph()
    p.space_after = Pt(7)
    r1 = p.add_run()
    r1.text = f"✓ {label}: "
    r1.font.bold = True
    r1.font.size = Pt(12)
    r1.font.color.rgb = ACCENT_GREEN
    r2 = p.add_run()
    r2.text = desc
    r2.font.size = Pt(11.5)
    r2.font.color.rgb = SLATE_TEXT

# Challenges & Mitigation Card
card_c = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.65), Inches(1.65), Inches(5.85), Inches(5.15))
card_c.fill.solid()
card_c.fill.fore_color.rgb = CARD_BG
card_c.line.color.rgb = BORDER_COLOR
card_c.line.width = Pt(1.5)

tf_c = card_c.text_frame
tf_c.word_wrap = True
tf_c.margin_left = Inches(0.25)
tf_c.margin_top = Inches(0.2)

p = tf_c.paragraphs[0]
p.text = "Challenges & Robust Mitigation Strategies"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = HEADER_BLUE
p.space_after = Pt(10)

challenges = [
    ("Asymmetric Flow State", "No reverse ACK/FIN packet in unidirectional taps.", "Mitigated via sliding-window timeout aggregation and statistical TCP sequence analysis."),
    ("Encrypted C2 / TLS 1.3", "Payload cannot be decrypted under privacy laws.", "Mitigated via metadata analytics: JA3/JA4 fingerprinting, Shannon Byte Entropy, and IAT variance."),
    ("High Traffic Volume", "Risk of buffer overflows during massive DDoS attacks.", "Mitigated by non-blocking asynchronous Python coroutines and selective client telemetry ring-buffering."),
    ("Concept Drift in Attacks", "Attackers altering beacon jitter and exfil rates.", "Mitigated through dual-engine consensus (Rule + ML) and 5-minute online baseline learning.")
]

for chal, risk, strat in challenges:
    p = tf_c.add_paragraph()
    p.space_after = Pt(6)
    r1 = p.add_run()
    r1.text = f"⚠ Challenge: {chal}\n"
    r1.font.bold = True
    r1.font.size = Pt(12)
    r1.font.color.rgb = ACCENT_AMBER
    r2 = p.add_run()
    r2.text = f"   • Risk: {risk}\n"
    r2.font.size = Pt(11)
    r2.font.color.rgb = MUTED_SLATE
    r3 = p.add_run()
    r3.text = f"   ✔ Solution: {strat}"
    r3.font.bold = True
    r3.font.size = Pt(11)
    r3.font.color.rgb = SLATE_TEXT

# ==========================================
# SLIDE 5: IMPACT AND BENEFITS
# ==========================================
slide5 = prs.slides.add_slide(blank_layout)
add_header(slide5, "IMPACT AND BENEFITS", 5)

sub_tb5 = slide5.shapes.add_textbox(Inches(0.8), Inches(1.15), Inches(11.7), Inches(0.4))
p_sub5 = sub_tb5.text_frame.paragraphs[0]
p_sub5.text = "❖ National Security Impact, Economic Advantages & Real-World Value"
p_sub5.font.size = Pt(18)
p_sub5.font.bold = True
p_sub5.font.color.rgb = BLUE_BAR

# 4 Quadrant Grid
qw = Inches(5.6)
qh = Inches(2.45)

quadrants = [
    ("🛡️ Impact on Target Audience & Defense", [
        "Intelligence Agencies (NTRO/RAW): Total stealth monitoring of foreign egress traffic without exposing tap presence.",
        "National Defense & Armed Forces: Air-gapped command and control network protection against advanced persistent threats (APTs).",
        "Critical Infrastructure (Power/Nuclear/Rail): Passive telemetry monitoring preventing catastrophic industrial cyber-sabotage."
    ], Inches(0.8), Inches(1.65)),
    ("💰 Economic & Operational Benefits", [
        "Eliminates Expensive Appliance Lock-In: Deploys on open-source stack and standard commodity servers, saving crores in license fees.",
        "Reduces SOC Analyst Fatigue by 70%: Correlated incidents and XAI explainability reduce false positive triage time from hours to seconds.",
        "Zero Downtime Deployment: 100% passive optical tap tap requires zero network interruption or scheduled maintenance downtime."
    ], Inches(6.65), Inches(1.65)),
    ("🌐 Social & Legal / Privacy Value", [
        "Strict Privacy Compliance: Analyzes packet metadata only; zero personal data decryption, adhering to DPDP Act 2023.",
        "Safeguards National Sovereignty: Protects public utility grids, banking infrastructure, and government data centers.",
        "Empowers Indigenous Cybersecurity: Built specifically under Make-in-India / Atmanirbhar Bharat vision for strategic cyber defense."
    ], Inches(0.8), Inches(4.3)),
    ("🚀 Scalability & Future Roadmap", [
        "Multi-Sensor Distributed Architecture: Scale from localized Wi-Fi taps to 100Gbps cross-country backbones.",
        "Continuous Model Evolution: Seamlessly retrains HistGradientBoosting weights on newly observed threat variants.",
        "Pre-Emptive Threat Interception: DTMC trajectory prediction identifies early reconnaissance before data exfiltration occurs."
    ], Inches(6.65), Inches(4.3))
]

for q_title, q_bullets, q_left, q_top in quadrants:
    q_card = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, q_left, q_top, qw, qh)
    q_card.fill.solid()
    q_card.fill.fore_color.rgb = CARD_BG
    q_card.line.color.rgb = BORDER_COLOR
    q_card.line.width = Pt(1.5)
    
    tf = q_card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.2)
    tf.margin_top = Inches(0.15)
    
    p = tf.paragraphs[0]
    p.text = q_title
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = HEADER_BLUE
    p.space_after = Pt(6)
    
    for b in q_bullets:
        pb = tf.add_paragraph()
        pb.space_after = Pt(4)
        parts = b.split(":", 1)
        r1 = pb.add_run()
        r1.text = "• " + parts[0] + (":" if len(parts) > 1 else "")
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = SLATE_TEXT
        if len(parts) > 1:
            r2 = pb.add_run()
            r2.text = parts[1]
            r2.font.size = Pt(10.5)
            r2.font.color.rgb = MUTED_SLATE

# Save presentation
output_path = os.path.abspath('CyberSentinel_SIH2026_Presentation.pptx')
prs.save(output_path)
print(f"Presentation generated successfully at: {output_path}")
